#!/usr/bin/env python
"""Build f1-presentation.pptx — a brushed-silver, editorial project deck."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots"
BRAND = ROOT / "frontend" / "src" / "assets" / "images" / "brand" / "f1.png"
OUT = ROOT / "f1-presentation.pptx"

W, H = 13.333, 7.5
PAGE = RGBColor(0xF6, 0xF7, 0xF8)
PANEL = RGBColor(0xFB, 0xFB, 0xFC)
INK = RGBColor(0x12, 0x15, 0x17)
BODY = RGBColor(0x3C, 0x42, 0x47)
MUTED = RGBColor(0x66, 0x6D, 0x74)
DIM = RGBColor(0x8D, 0x94, 0x9B)
LINE = RGBColor(0xD5, 0xDA, 0xDE)
ACCENT = RGBColor(0xE1, 0x06, 0x00)
ACCENT_SOFT = RGBColor(0xFB, 0xE8, 0xE7)

FONT = "Segoe UI"
FONT_DISPLAY = "Segoe UI Semibold"
FONT_MONO = "Consolas"


# ---------------------------------------------------------------------------
# primitives
# ---------------------------------------------------------------------------
def no_line(shape) -> None:
    shape.line.fill.background()


def rect(slide, x, y, w, h, color, line=None, radius=None, line_width=0.75):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    else:
        no_line(shape)
    shape.shadow.inherit = False
    if radius and shape.adjustments:
        shape.adjustments[0] = radius
    shape.text_frame.word_wrap = True
    shape.text_frame.text = ""
    return shape


def text(
    slide,
    x,
    y,
    w,
    h,
    value,
    size=16,
    color=INK,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    spacing=1.0,
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    lines = value if isinstance(value, list) else [value]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = spacing
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return box


def kicker(slide, x, y, label, color=ACCENT, rule=True):
    text(slide, x, y, 4.4, 0.24, label.upper(), size=8.5, color=color, bold=True, font=FONT_MONO)
    if rule:
        rect(slide, x, y + 0.23, 0.46, 0.028, color)


def background(slide, number: int, total: int, section: str | None = None) -> None:
    rect(slide, 0, 0, W, H, PAGE)

    # Header rail.
    rect(slide, 0.62, 0.46, 0.055, 0.26, ACCENT)
    if section:
        text(slide, 0.78, 0.48, 8.4, 0.26, section.upper(), size=9, color=MUTED, bold=True, font=FONT_MONO)

    # Footer rail.
    rect(slide, 0.62, 6.93, 12.1, 0.012, LINE)
    text(
        slide,
        0.62,
        7.04,
        7.0,
        0.25,
        "PADDOCK  ·  FORMULA 1 ARCHIVE",
        size=8,
        color=DIM,
        font=FONT_MONO,
    )
    text(
        slide,
        W - 1.55,
        7.04,
        0.9,
        0.25,
        f"{number:02d} / {total:02d}",
        size=8,
        color=DIM,
        align=PP_ALIGN.RIGHT,
        font=FONT_MONO,
    )


def heading(slide, title, subtitle=None, y=0.82):
    text(slide, 0.62, y, 11.9, 0.66, title, size=29, font=FONT_DISPLAY, bold=True, spacing=0.95)
    if subtitle:
        text(slide, 0.64, y + 0.7, 10.9, 0.5, subtitle, size=12, color=MUTED, spacing=1.18)


def image(slide, name, x, y, w=None, h=None, border=LINE):
    path = SHOTS / name
    if not path.exists():
        return None
    kwargs = {}
    if w:
        kwargs["width"] = Inches(w)
    if h:
        kwargs["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)
    pic.line.color.rgb = border
    pic.line.width = Pt(0.75)
    return pic


def card(slide, x, y, w, h, accent=None, fill=PANEL, line=LINE, radius=0.035):
    shape = rect(slide, x, y, w, h, fill, line=line, radius=radius)
    if accent:
        rect(slide, x, y, 0.045, h, accent)
    return shape


def stat(slide, x, y, w, value, label, color=INK):
    text(slide, x, y, w, 0.52, value, size=27, font=FONT_DISPLAY, bold=True, color=color)
    text(slide, x, y + 0.55, w, 0.28, label.upper(), size=8, color=MUTED, bold=True, font=FONT_MONO)


def notes(slide, value: str) -> None:
    slide.notes_slide.notes_text_frame.text = value


# ---------------------------------------------------------------------------
# slides
# ---------------------------------------------------------------------------
def title_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 1, total, "FET project")

    if BRAND.exists():
        slide.shapes.add_picture(str(BRAND), Inches(11.98), Inches(0.34), width=Inches(0.78))

    kicker(slide, 0.62, 1.16, "2000 → 2026")
    text(slide, 0.6, 1.52, 6.5, 1.15, "PADDOCK", size=62, font=FONT_DISPLAY, bold=True, spacing=0.9)
    text(
        slide,
        0.64,
        2.7,
        5.4,
        1.1,
        ["Every season since 2000.", "Every lap since 2023."],
        size=20,
        color=BODY,
        spacing=1.12,
    )

    rect(slide, 0.64, 3.94, 1.3, 0.028, ACCENT)
    text(
        slide,
        0.64,
        4.2,
        5.2,
        1.2,
        "A brushed-silver Formula 1 archive built with Angular 21 and Material 3. "
        "The complete 2000–2026 dataset meets OpenF1 session replay, live timing and telemetry.",
        size=11.5,
        color=MUTED,
        spacing=1.38,
    )

    text(slide, 0.64, 5.82, 4.6, 0.3, "JUGAL PATEL · KRISH PATEL · PAAVAN SHARMA", size=8.4, color=INK, bold=True, font=FONT_MONO)
    text(slide, 0.64, 6.1, 4.6, 0.3, "24162581032 · 24162581017 · 24162581052", size=8, color=DIM, font=FONT_MONO)
    text(slide, 0.64, 6.38, 4.6, 0.3, "FET · SEMESTER 5", size=8, color=DIM, font=FONT_MONO)

    if image(slide, "home.png", 6.62, 1.12, w=6.25):
        text(
            slide,
            6.64,
            5.82,
            6.2,
            0.3,
            "HOME  ·  MP4/4 HERO  ·  2026 TOP FIVE  ·  526 RACES",
            size=8,
            color=MUTED,
            bold=True,
            font=FONT_MONO,
        )

    notes(
        slide,
        "Opening: Paddock is a complete redesign of the F1 stats app. The archive now runs to 26 October 2026, "
        "the interface is a single brushed-silver editorial theme, and OpenF1 supplies everything from 2023 onward.",
    )


def brief_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 2, total, "The brief")
    heading(
        slide,
        "The old app knew a lot. It showed very little.",
        "A full data archive sat behind a front end that was hard to read, hard to navigate and out of date.",
    )

    cards = [
        ("01", "Fragmented experience", "Nine pages with different visual rules, inconsistent spacing and no clear path from a season to a race to a driver."),
        ("02", "Data without context", "Raw result tables. No career totals, no championship picture, no comparison tools, no telemetry."),
        ("03", "Stuck in the past", "The archive stopped at 2024 and the current generation of cars had no live or replay layer."),
    ]
    for index, (number, title, body) in enumerate(cards):
        x = 0.62 + index * 4.08
        card(slide, x, 2.3, 3.8, 2.35, accent=ACCENT if index == 0 else None)
        text(slide, x + 0.28, 2.52, 0.6, 0.3, number, size=10, color=ACCENT, bold=True, font=FONT_MONO)
        text(slide, x + 0.28, 2.86, 3.2, 0.5, title, size=15.5, font=FONT_DISPLAY, bold=True)
        text(slide, x + 0.28, 3.46, 3.22, 1.0, body, size=10, color=MUTED, spacing=1.28)

    rect(slide, 0.62, 5.12, 12.1, 1.15, PANEL, line=LINE, radius=0.03)
    rect(slide, 0.62, 5.12, 0.045, 1.15, ACCENT)
    text(slide, 0.92, 5.31, 2.0, 0.3, "THE GOAL", size=8.5, color=ACCENT, bold=True, font=FONT_MONO)
    text(
        slide,
        2.35,
        5.26,
        9.95,
        0.85,
        "One coherent product: a calm, fast, 2000-onward archive with a real live layer — shipped end to end in one pass.",
        size=14.5,
        color=INK,
        spacing=1.18,
    )

    notes(
        slide,
        "Frame the redesign: the archive was already rich, but the product was not. The goal was cohesion, reach "
        "and a quieter visual system, not just a reskin.",
    )


def glance_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 3, total, "Product")
    heading(slide, "One paddock, eleven screens.", "Every route shares one silver system, and every table drills into a profile or a race.")

    stats = [
        ("27", "seasons bundled"),
        ("526", "archived races"),
        ("10,954", "race results"),
        ("11", "screens"),
    ]
    for index, (value, label) in enumerate(stats):
        x = 0.62 + index * 3.12
        card(slide, x, 2.2, 2.9, 1.2)
        stat(slide, x + 0.25, 2.35, 2.4, value, label)

    left = [
        ("/", "Home", "MP4/4 hero, 2026 board, latest and next race"),
        ("/standings", "Standings", "Drivers' and constructors' tables, leader logic"),
        ("/drivers", "Drivers", "Searchable grid with career profile dialogs"),
        ("/constructors", "Teams", "Real logos or typographic crests, title years"),
        ("/circuits", "Circuits", "39 layouts, searchable map and detail panel"),
        ("/races", "Races", "Season calendar, podiums and year switcher"),
    ]
    right = [
        ("/races/:id", "Race detail", "Classification, qualifying and season context"),
        ("/live", "Live timing", "OpenF1 replay, timing tower and telemetry"),
        ("/history", "History", "Wikipedia-style eras, rivalries and records"),
        ("/rivalry", "Rivalry", "Any two drivers, career totals and shared seasons"),
        ("/credits", "Credits", "Every image, its source and its licence"),
    ]

    for column, rows, x0 in ((0, left, 0.62), (1, right, 6.85)):
        for index, (route, name, body) in enumerate(rows):
            y = 3.72 + index * 0.54
            text(slide, x0, y, 1.55, 0.24, route, size=7.8, color=MUTED, font=FONT_MONO)
            text(slide, x0 + 1.62, y - 0.03, 1.5, 0.28, name, size=10, color=INK, bold=True)
            text(slide, x0 + 1.62, y + 0.23, 4.1, 0.26, body, size=8.3, color=MUTED)

    notes(
        slide,
        "Scope: eleven routes, including the new credits page. The archive holds 27 seasons, 526 races and "
        "10,954 results after the 2026 reconciliation.",
    )


def design_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 4, total, "Design language")
    heading(
        slide,
        "Brushed silver, one red signal.",
        "A flat editorial system, not a dark dashboard. One theme, hairline rules, and type doing the work.",
    )

    if image(slide, "circuits.png", 7.15, 1.92, w=5.55):
        pass

    items = [
        ("One theme, no toggle", "Silver is the only surface. Dark mode, glass and glow were deleted, not restyled."),
        ("Flat panels", "One .panel primitive: matte surface, hairline border, 2–4 px radius, one soft shadow."),
        ("Two typefaces", "Titillium Web for UI and headlines, Source Serif 4 for the long-form history."),
        ("Red as a signal", "F1 mark, active nav rule, P1 and one primary action per view. Never a background."),
        ("Maps as field", "Thirty distinct black outlines at 20% ink cover the full canvas; ten on mobile, none under reduced motion."),
    ]
    for index, (title, body) in enumerate(items):
        y = 1.95 + index * 0.94
        card(slide, 0.62, y, 6.15, 0.8)
        text(slide, 0.85, y + 0.11, 5.6, 0.26, title, size=11, bold=True)
        text(slide, 0.85, y + 0.38, 5.65, 0.36, body, size=8.8, color=MUTED, spacing=1.12)

    notes(
        slide,
        "Design: silver surfaces, 2–4 px radii and hairline rules. Red appears only on the F1 mark, the active nav "
        "rule, P1 and one primary action. The circuits page shows the flat panel and layout-image treatment.",
    )


def architecture_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 5, total, "Architecture")
    heading(
        slide,
        "A small front end with a clear spine.",
        "Standalone Angular components, signals for state, and two services that own all data access.",
    )

    layers = [
        ("SHELL", "app.ts · navbar · router outlet · floating circuit layer", ACCENT),
        ("PAGES", "home · standings · drivers · teams · circuits · races · race detail · live · history · rivalry · credits", INK),
        ("SHARED UI", "panel · page header · stat tile · driver avatar · team badge · track background · charts", INK),
        ("CORE", "CSVDataService · OpenF1Service · models · config · utils", INK),
    ]
    for index, (label, body, color) in enumerate(layers):
        y = 2.12 + index * 1.02
        rect(slide, 0.62, y, 7.4, 0.86, PANEL, line=LINE, radius=0.03)
        rect(slide, 0.62, y, 0.045, 0.86, color)
        text(slide, 0.86, y + 0.12, 1.6, 0.26, label, size=8.5, color=ACCENT if index == 0 else MUTED, bold=True, font=FONT_MONO)
        text(slide, 2.35, y + 0.1, 5.45, 0.58, body, size=9.3, color=INK, spacing=1.14)

    card(slide, 8.32, 2.12, 4.38, 3.86)
    text(slide, 8.58, 2.32, 3.9, 0.3, "KEY DECISIONS", size=8.5, color=MUTED, bold=True, font=FONT_MONO)
    decisions = [
        ("Lazy routes", "Every page except Home is a separate chunk."),
        ("Indexed CSVs", "Parse once, derive careers and standings in memory."),
        ("shareReplay", "One HTTP request per archive file and OpenF1 call."),
        ("Rate limited", "3 req/s token bucket protects the OpenF1 free tier."),
        ("Typed end to end", "Strict TypeScript models for archive and OpenF1."),
    ]
    for index, (title, body) in enumerate(decisions):
        y = 2.72 + index * 0.62
        text(slide, 8.58, y, 3.8, 0.28, "— " + title, size=10, color=INK, bold=True)
        text(slide, 8.86, y + 0.28, 3.55, 0.3, body, size=8.5, color=MUTED)

    notes(
        slide,
        "Architecture: the shell holds navigation and the floating map layer; pages are lazy; shared components "
        "keep the visual language consistent; two services own data. No backend is required at runtime.",
    )


def data_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 6, total, "Data")
    heading(
        slide,
        "Two sources. One continuous timeline.",
        "The archive gives depth and precision; OpenF1 gives recency and live detail.",
    )

    # Archive block.
    card(slide, 0.62, 2.12, 5.85, 3.45, accent=ACCENT)
    kicker(slide, 0.92, 2.3, "2000 → 2026")
    text(slide, 0.92, 2.66, 4.9, 0.4, "Bundled archive", size=19, font=FONT_DISPLAY, bold=True)
    text(slide, 0.92, 3.1, 5.2, 0.5, "Ergast-style CSVs, pruned in the browser build.", size=10, color=MUTED)
    facts = [
        "130 drivers  ·  40 teams  ·  39 circuits",
        "526 races  ·  10,954 results  ·  9,968 qualifying",
        "Final or current standings for all 27 seasons",
        "~1.3 MB total, loaded locally, indexed once",
    ]
    for index, fact in enumerate(facts):
        text(slide, 0.92, 3.68 + index * 0.34, 5.2, 0.3, "·  " + fact, size=9.4, color=INK)

    # OpenF1 block.
    card(slide, 6.78, 2.12, 5.92, 3.45, accent=INK)
    kicker(slide, 7.08, 2.3, "2023 → TODAY", color=INK)
    text(slide, 7.08, 2.66, 5.0, 0.4, "OpenF1 API", size=19, font=FONT_DISPLAY, bold=True)
    text(slide, 7.08, 3.1, 5.2, 0.5, "Session replay, timing and telemetry.", size=10, color=MUTED)
    facts = [
        "Sessions  ·  meetings  ·  drivers  ·  positions",
        "Intervals  ·  laps  ·  stints  ·  pit stops",
        "Weather  ·  race control  ·  car data",
        "3 req/s cap with cache + token bucket",
    ]
    for index, fact in enumerate(facts):
        text(slide, 7.08, 3.68 + index * 0.34, 5.3, 0.3, "·  " + fact, size=9.4, color=INK)

    text(slide, 6.46, 3.5, 0.3, 0.4, "→", size=18, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

    rect(slide, 0.62, 5.85, 12.08, 0.78, ACCENT_SOFT, line=LINE, radius=0.03)
    text(
        slide,
        0.92,
        5.99,
        11.5,
        0.5,
        "2026 is reconciled to round 18 of 23 (26 October 2026). An unfinished season shows a leader — never a champion.",
        size=10.5,
        color=INK,
        italic=True,
    )

    notes(
        slide,
        "Data: the archive is trimmed at build time and indexed at runtime. 2025 and 2026 come from Jolpica, "
        "including the new Audi, Cadillac and Madring entries. OpenF1 is the live layer.",
    )


def screens_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 7, total, "Screens")
    heading(
        slide,
        "The season at a glance.",
        "Home and standings share one hierarchy: context first, numbers second, detail one click away.",
    )

    if image(slide, "home.png", 0.62, 1.82, w=6.0):
        pass
    if image(slide, "standings.png", 6.85, 1.82, w=5.85):
        pass
    text(slide, 0.62, 6.05, 5.9, 0.3, "HOME", size=8, color=ACCENT, bold=True, font=FONT_MONO)
    text(slide, 0.62, 6.3, 5.9, 0.5, "MP4/4 hero, 2026 top five, archive totals, latest and next race.", size=9.3, color=MUTED)
    text(slide, 6.85, 6.05, 5.8, 0.3, "STANDINGS", size=8, color=MUTED, bold=True, font=FONT_MONO)
    text(slide, 6.85, 6.3, 5.8, 0.5, "Drivers' and constructors' tables, with a leader card for an unfinished season.", size=9.3, color=MUTED)

    notes(
        slide,
        "Two flagship screens. Home answers 'what is happening'; standings answers 'who is ahead and by how much'. "
        "Both use the same panel, type scale and red-on-silver hierarchy.",
    )


def race_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 8, total, "Race detail")
    heading(slide, "Every race gets a proper story.", "A classification page turns a row in a CSV into a readable race report.")

    if image(slide, "race-detail.png", 0.62, 1.85, w=7.0):
        pass

    points = [
        ("Winner hero", "A 2 px team-colour rule, podium, race time and grid drop."),
        ("Fastest lap", "Driver, lap time and top speed in one card."),
        ("Classification", "Grid delta, laps, time/status, points and FL tag."),
        ("Qualifying", "Q1 / Q2 / Q3 tab with the full session order."),
        ("Context", "Circuit, season stats and previous/next round."),
    ]
    for index, (title, body) in enumerate(points):
        y = 2.06 + index * 0.92
        card(slide, 8.55, y, 4.15, 0.78)
        text(slide, 8.75, y + 0.1, 3.7, 0.28, title, size=10.5, bold=True)
        text(slide, 8.75, y + 0.38, 3.75, 0.34, body, size=8.6, color=MUTED, spacing=1.1)

    notes(
        slide,
        "Race detail links the result back to the season, the circuit and the driver, and exposes qualifying that "
        "used to be invisible. Team colour is a 2 px rule, not a background.",
    )


def live_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 9, total, "Live timing")
    heading(
        slide,
        "Replay today. Poll live when you have a token.",
        "OpenF1 data drives a timing tower, race control, weather and a windowed telemetry trace.",
    )

    if image(slide, "live.png", 0.62, 1.85, w=7.0):
        pass

    labels = [
        ("Timing tower", "Position, gap, interval, laps, last/best, tyre age, pit stops."),
        ("Race control", "Flag-coloured message feed with lap and timestamp."),
        ("Telemetry", "Speed, gear, DRS and gear changes for one lap only."),
        ("Safe by default", "Polling off until a bearer token is stored locally."),
    ]
    for index, (title, body) in enumerate(labels):
        y = 2.06 + index * 1.12
        card(slide, 8.55, y, 4.15, 0.96, accent=ACCENT if index == 0 else None)
        text(slide, 8.76, y + 0.13, 3.7, 0.28, title, size=10.5, bold=True)
        text(slide, 8.76, y + 0.44, 3.75, 0.42, body, size=8.6, color=MUTED, spacing=1.12)

    notes(
        slide,
        "Live is replay-first because OpenF1's free tier is historical. A paid token can be pasted into the page to "
        "enable 15-second polling; it stays in the browser. Telemetry is fetched per driver, per lap.",
    )


def analysis_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 10, total, "Analysis")
    heading(
        slide,
        "The archive becomes an argument.",
        "Records, eras and head-to-head comparisons are computed from the same rows that power the tables.",
    )

    if image(slide, "rivalry.png", 0.62, 1.82, w=5.7):
        pass
    if image(slide, "history.png", 6.78, 1.82, w=5.7):
        pass

    blocks = [
        ("Rivalry engine", "Any two of 130 drivers. Career totals, points trajectory, shared seasons and a head-to-head score.", 0.62),
        ("History article", "Wikipedia-style lead, infobox, five eras, four rivalries, champions table and records.", 6.78),
    ]
    for title, body, x in blocks:
        rect(slide, x, 5.78, 5.7, 0.92, PANEL, line=LINE, radius=0.03)
        rect(slide, x, 5.78, 0.045, 0.92, ACCENT if x < 1 else INK)
        text(slide, x + 0.22, 5.9, 5.2, 0.28, title, size=10, bold=True)
        text(slide, x + 0.22, 6.18, 5.2, 0.42, body, size=8.4, color=MUTED, spacing=1.1)

    notes(
        slide,
        "The analysis layer is where the archive earns its keep: comparisons, eras and records are derived live, "
        "so they stay correct when the data changes. /timeline now redirects to /history.",
    )


def imagery_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 11, total, "Imagery")
    heading(
        slide,
        "Every team and every circuit has an image.",
        "Real marks where they can be licensed, designed fallbacks where they cannot — and credit for all of it.",
    )

    if image(slide, "teams.png", 0.62, 1.82, w=5.7):
        pass
    if image(slide, "credits.png", 6.78, 1.82, w=5.7):
        pass

    rows = [
        ("35 of 40", "team logos", "Five defunct teams render a typographic crest."),
        ("39 of 39", "circuit maps", "Every venue resolves to its official layout image."),
        ("28", "driver portraits", "OpenF1 headshots; earlier drivers keep monograms."),
        ("9", "archive photos", "MP4/4, Hunt–Lauda, Senna and more, credited in-app."),
    ]
    for index, (value, label, body) in enumerate(rows):
        x = 0.62 + index * 3.05
        text(slide, x, 5.92, 2.9, 0.34, value, size=15, font=FONT_DISPLAY, bold=True, color=ACCENT if index == 0 else INK)
        text(slide, x, 6.22, 2.9, 0.24, label.upper(), size=7.5, color=MUTED, bold=True, font=FONT_MONO)
        text(slide, x, 6.46, 2.75, 0.34, body, size=7.8, color=MUTED, spacing=1.1)

    notes(
        slide,
        "The generated manifest lists only files that exist, so a missing asset never becomes a broken image. "
        "Five historic teams use crests, and every image is attributed on /credits with a trademark notice.",
    )


def performance_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 12, total, "Quality")
    heading(
        slide,
        "Fast, accessible, and honest about limits.",
        "The build is lean, the interface is keyboard-friendly, and the app says when data is unavailable.",
    )

    stats = [
        ("580 kB", "initial raw bundle"),
        ("154 kB", "initial transfer"),
        ("10", "lazy routes"),
        ("3 req/s", "OpenF1 budget"),
    ]
    for index, (value, label) in enumerate(stats):
        x = 0.62 + index * 3.12
        card(slide, x, 2.18, 2.9, 1.2)
        stat(slide, x + 0.25, 2.33, 2.4, value, label)

    columns = [
        ("PERFORMANCE", [
            "Route-level code splitting",
            "Archived CSVs indexed once",
            "shareReplay on every request",
            "Budgets enforced in angular.json",
        ]),
        ("ACCESSIBILITY", [
            "Semantic headings and captions",
            "Labelled controls and tab states",
            "Visible focus and keyboard nav",
            "Reduced motion / transparency",
        ]),
        ("RESILIENCE", [
            "Skeletons instead of blank screens",
            "Designed empty and error states",
            "Crest fallback for missing logos",
            "Flag fallback for missing maps",
        ]),
    ]
    for index, (title, items) in enumerate(columns):
        x = 0.62 + index * 4.08
        card(slide, x, 3.8, 3.78, 2.5, accent=ACCENT if index == 0 else None)
        text(slide, x + 0.25, 3.97, 3.3, 0.3, title, size=8.5, color=ACCENT if index == 0 else MUTED, bold=True, font=FONT_MONO)
        for row, item in enumerate(items):
            text(slide, x + 0.25, 4.4 + row * 0.42, 3.35, 0.3, "·  " + item, size=9.2, color=INK)

    notes(
        slide,
        "Quality: the production initial bundle is 580 kB raw, or 154 kB over the wire. Every page except Home is "
        "lazy. Third-party glyphs and missing images always have a designed fallback.",
    )


def run_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 13, total, "Delivery")
    heading(slide, "Clone. Install. Run.", "No database, no API key and no build step for the bundled archive.")

    card(slide, 0.62, 2.08, 6.55, 4.0, accent=ACCENT)
    text(slide, 0.92, 2.26, 5.9, 0.3, "QUICK START", size=8.5, color=ACCENT, bold=True, font=FONT_MONO)
    commands = [
        ("cd frontend", INK),
        ("npm install", INK),
        ("npm start", INK),
        ("", INK),
        ("# production", DIM),
        ("npm run build", INK),
        ("", INK),
        ("# full data refresh", DIM),
        ("npm run data:build", INK),
        ("npm run assets:fetch", INK),
        ("npm run outlines:build", INK),
    ]
    for index, (command, color) in enumerate(commands):
        text(slide, 0.92, 2.66 + index * 0.3, 5.8, 0.28, command or " ", size=9.8, color=color, font=FONT_MONO)

    card(slide, 7.45, 2.08, 5.25, 4.0)
    text(slide, 7.75, 2.26, 4.6, 0.3, "WHAT SHIPS", size=8.5, color=MUTED, bold=True, font=FONT_MONO)
    shipped = [
        ("App", "11 routes · silver editorial · responsive"),
        ("Archive", "2000–2026 · 1.3 MB · pruned CSVs"),
        ("Live", "OpenF1 replay + optional token polling"),
        ("Images", "35 logos · 39 maps · 28 portraits · 9 photos"),
        ("Docs", "README.md · README.docx · README.pdf"),
        ("Deck", "f1-presentation.pptx · .pdf"),
    ]
    for index, (title, body) in enumerate(shipped):
        y = 2.7 + index * 0.55
        text(slide, 7.75, y, 1.05, 0.28, title, size=9.5, color=INK, bold=True)
        text(slide, 8.85, y, 3.7, 0.28, body, size=8.6, color=MUTED)

    notes(
        slide,
        "Delivery: the front end runs on its own. The prune, sync, asset and outline scripts are reproducible, and "
        "the docs and deck ship with the repo.",
    )


def limits_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 14, total, "Next")
    heading(slide, "Known edges. Clear next moves.", "The product is complete for this scope, and the boundaries are deliberate.")

    left = [
        ("2026 is still running", "18 of 23 rounds are complete. The app labels the current leader and lists no 2026 champion."),
        ("Portraits start in 2023", "OpenF1 only lists modern headshots. Earlier drivers keep the designed monogram."),
        ("Live needs a token", "OpenF1's free tier is historical. Real-time polling requires a paid subscription."),
        ("Some marks cannot be licensed", "Five defunct teams use a typographic crest instead of a logo."),
    ]
    right = [
        ("Sprint + lap data", "Add sprint results and lap-by-lap archive data for deeper race analysis."),
        ("Favourites + sharing", "Persist favourite drivers and shareable comparison links."),
        ("Caching layer", "A small server cache for OpenF1 would make live pages instant for everyone."),
        ("More era photography", "Extend the credited archive photo set beyond the nine key moments."),
    ]
    for column, (title, items, accent) in enumerate([
        ("BOUNDARIES", left, ACCENT),
        ("NEXT", right, INK),
    ]):
        x = 0.62 + column * 6.23
        card(slide, x, 2.12, 5.85, 4.1, accent=accent)
        text(slide, x + 0.28, 2.3, 5.2, 0.3, title, size=8.5, color=ACCENT if column == 0 else MUTED, bold=True, font=FONT_MONO)
        for index, (heading_text, body) in enumerate(items):
            y = 2.74 + index * 0.82
            text(slide, x + 0.28, y, 5.2, 0.3, heading_text, size=10.8, bold=True)
            text(slide, x + 0.28, y + 0.3, 5.15, 0.44, body, size=8.7, color=MUTED, spacing=1.14)

    notes(
        slide,
        "Be explicit about limits. They are not failures: they reflect data contracts. The next moves are natural "
        "extensions, not fixes.",
    )


def close_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide, 15, total, "Closing")

    kicker(slide, 0.62, 1.5, "THANK YOU")
    text(slide, 0.6, 1.9, 8.4, 1.2, "The paddock is open.", size=52, font=FONT_DISPLAY, bold=True)
    text(
        slide,
        0.64,
        3.16,
        7.0,
        0.9,
        "Questions, laps, or a debate about the 2021 finale — all welcome.",
        size=15,
        color=MUTED,
    )

    rect(slide, 0.64, 4.4, 1.3, 0.028, ACCENT)

    text(slide, 0.64, 4.7, 5.2, 0.35, "JUGAL PATEL · KRISH PATEL · PAAVAN SHARMA", size=10.5, color=INK, bold=True, font=FONT_MONO)
    text(slide, 0.64, 5.06, 5.2, 0.3, "24162581032 · 24162581017 · 24162581052 · FET", size=8.5, color=DIM, font=FONT_MONO)
    text(slide, 0.64, 5.64, 5.2, 0.3, "localhost:4200", size=10, color=INK, font=FONT_MONO)
    text(slide, 6.4, 5.64, 5.2, 0.3, "openf1.org  ·  angular.dev", size=10, color=MUTED, font=FONT_MONO)

    if BRAND.exists():
        slide.shapes.add_picture(str(BRAND), Inches(9.2), Inches(2.35), width=Inches(3.1))

    notes(
        slide,
        "Close: thank you. Invite questions about the silver design system, the 2026 data reconciliation or the "
        "OpenF1 integration.",
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    total = 15

    title_slide(prs, total)
    brief_slide(prs, total)
    glance_slide(prs, total)
    design_slide(prs, total)
    architecture_slide(prs, total)
    data_slide(prs, total)
    screens_slide(prs, total)
    race_slide(prs, total)
    live_slide(prs, total)
    analysis_slide(prs, total)
    imagery_slide(prs, total)
    performance_slide(prs, total)
    run_slide(prs, total)
    limits_slide(prs, total)
    close_slide(prs, total)

    prs.core_properties.title = "Paddock — Formula 1 archive"
    prs.core_properties.author = "Jugal Patel, Krish Patel, Paavan Sharma"
    prs.core_properties.subject = "FET project presentation"
    prs.core_properties.comments = "Angular 21 + Material 3 + Jolpica + OpenF1"

    prs.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT)} ({OUT.stat().st_size / 1024:.0f} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
